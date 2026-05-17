import subprocess
import re
import asyncio
import edge_tts
import os
from pptx import Presentation

# Configuration
PHP_MINER_PATH = "/var/www/html/moodle_test/local/chartplugin/test_seed.php"
AUDIO_DIR = "/var/www/html/moodle_test/local/chartplugin/media_assets/audio"
VOICE = "en-GB-SoniaNeural"

async def generate_voiceover(text, filename):
    """Q918: Debonair Vocal Branding"""
    filepath = os.path.join(AUDIO_DIR, filename)
    
    if os.path.exists(filepath):
        print(f"--- Skipping Voice: {filename} (Already exists) ---")
        return

    print(f"--- Synthesizing DEBONAIR Voice: {filename} ---")
    try:
        communicate = edge_tts.Communicate(text, VOICE, rate="-5%", pitch="-2Hz")
        await communicate.save(filepath)
    except Exception as e:
        print(f"ERROR: {e}")

def harvest_from_moodle():
    """Q907: CLI Handshake - Handles the 'Content:' label layout"""
    print(f"--- Running PHP Miner locally: {PHP_MINER_PATH} ---")
    try:
        result = subprocess.run(['php', PHP_MINER_PATH], capture_output=True, text=True)
        content = result.stdout
        if not content:
            print("ERROR: No content received from PHP.")
            return []
        
        # 3-Group Capture Pattern: 1=Title, 2=Slide Header, 3=Raw Narrative Package Block
        pattern = r"Found:\s*(.*?)\n\s*Content:\s*\[SLIDE:\s*(.*?)\s*\]\s*.*?\[NARRATIVE:\s*(.*?)\s*\]"
        return re.findall(pattern, content, re.DOTALL | re.IGNORECASE)
    except Exception as e:
        print(f"CRITICAL ERROR during harvest: {e}")
        return []
    
async def main():
    lessons = harvest_from_moodle()
    print(f"Match Count: {len(lessons)}")
    
    if not lessons:
        print("!!! NO READY LESSONS FOUND. CHECK MOODLE TAGS. !!!")
        return

    prs = Presentation()
    
    # --- TITLE SLIDE (Index 1) ---
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Debonair AI: Moodle API Series"
    slide.placeholders[1].text = f"Automated Generation: {len(lessons)} Lessons"
    
    # --- CONTENT SLIDES (Starting at Slide Index 2) ---
    for idx, (title, slide_header, raw_narrative_block) in enumerate(lessons, start=2):
        clean_title = title.strip()
        clean_header = slide_header.strip()
        narrative_content = raw_narrative_block.strip()
        
        print(f"Processing Slide {idx}: {clean_title}")
        
        # 1. Structural Data Isolation: Split Bullet Points block from the Voiceover Script
        parts = re.split(r"Voiceover\s*Script:", narrative_content, flags=re.IGNORECASE)
        
        bullet_segment = parts[0].replace("Slide Bullet Points:", "").strip()
        voiceover_text = parts[1].strip() if len(parts) > 1 else narrative_content
        
        # 2. Create Slide using Title and Content Layout
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = clean_header
        
        # 3. Add Voiceover Script to Notes viewport frame
        slide.notes_slide.notes_text_frame.text = voiceover_text
        
        # 4. Target Content Box Placeholder and Populate Slide Body Canvas
        content_placeholder = slide.placeholders[1]
        tf = content_placeholder.text_frame
        tf.clear()  # Clear standard template placeholder text
        
        # Split text data block line-by-line to parse out individual bullets
        lines = [line.strip() for line in bullet_segment.split("\n") if line.strip()]
        
        first_bullet = True
        for line in lines:
            # Strip loose bullet artifact signs remaining from plain-text layout formatting
            clean_line = re.sub(r"^[•\-\*]\s*", "", line).strip()
            if not clean_line:
                continue
                
            if first_bullet:
                p = tf.paragraphs[0]
                p.text = clean_line
                first_bullet = False
            else:
                p = tf.add_paragraph()
                p.text = clean_line
                p.level = 0  # Anchor bullet list at default base root indent
        
        # 5. Generate Track using unique Slide Sequence names to avoid offsets
        safe_name = "".join([c for c in clean_title if c.isalnum() or c in (' ', '_')]).replace(' ', '_')
        audio_filename = f"slide_{idx}_{safe_name}.mp3"
        audio_filepath = os.path.join(AUDIO_DIR, audio_filename)
        
        await generate_voiceover(voiceover_text, audio_filename)
        
        # 6. Native PowerPoint Media Link Insertion
        if os.path.exists(audio_filepath):
            try:
                # Embeds a native playback shape directly inside the slide container layout
                slide.shapes.add_movie(
                    audio_filepath, 
                    left=0, top=0, 
                    width=500000, height=500000, 
                    poster_frame_image=None
                )
                print(f"--> Automatically embedded {audio_filename} onto Slide {idx}")
            except Exception as media_err:
                print(f"Media embedding warning on slide {idx}: {media_err}")
        
        print(f"SUCCESS: Assets ready for {clean_title}")

    # Final Save
    output_pptx = "Moodle_Course_Bulk_Draft.pptx"
    prs.save(output_pptx)
    print(f"--- {output_pptx} CREATED SUCCESSFULLY ---")

if __name__ == "__main__":
    if not os.path.exists(AUDIO_DIR):
        os.makedirs(AUDIO_DIR, exist_ok=True)
    
    asyncio.run(main())